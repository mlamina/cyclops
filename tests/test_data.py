"""What ``Project Data.xlsx`` holds, what survives a save, and how a value is found again.

Imports only ``cyclops.projects.data`` on purpose - that module never touches the disk and needs
no key, no camera and no card, so this runs anywhere in well under a second. Everything about
*where* the file lives is ``store.py``'s business and is checked on the Pi, not here.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from cyclops.projects import data

# ------------------------------------------------------------------ building books to judge


def book(*rows):
    """A book holding ``(tab, key, value[, note])`` tuples, in the order given."""
    made = data.Book()
    for row in rows:
        made.put(*row)
    return made


def keys(rows):
    return [row.key for row in rows]


VICE = (
    ("Torque specs", "Caliper bolt torque", "25 Nm", "dry thread"),
    ("Torque specs", "Carrier bolt torque", "110 Nm"),
    ("Dimensions", "Jaw width", "125 mm"),
)


# ------------------------------------------------------------------ the file


def test_round_trip_keeps_every_row_and_its_tab():
    back = data.load(data.dump(book(*VICE)))
    assert back.tabs() == {"Torque specs": 2, "Dimensions": 1}
    assert [(r.tab, r.key, r.value, r.note) for r in back.rows()] == [
        ("Torque specs", "Caliper bolt torque", "25 Nm", "dry thread"),
        ("Torque specs", "Carrier bolt torque", "110 Nm", ""),
        ("Dimensions", "Jaw width", "125 mm", ""),
    ]


def test_an_empty_book_is_still_a_workbook_that_opens():
    """openpyxl refuses to save a workbook with no sheets, and Excel refuses to open one."""
    assert load_workbook(io.BytesIO(data.dump(data.Book()))).sheetnames


def test_a_missing_or_unreadable_file_is_an_empty_book_not_an_error():
    assert data.load(None).rows() == []
    assert data.load(b"").rows() == []
    assert data.load(b"this is not a zip file").rows() == []


def test_the_header_is_written_and_not_read_back_as_a_value():
    sheet = load_workbook(io.BytesIO(data.dump(book(*VICE))))["Dimensions"]
    assert [c.value for c in sheet[1]] == list(data.HEADERS)
    assert sheet[1][0].font.bold


def test_a_sheet_someone_stripped_the_header_off_still_reads_every_row():
    hand = Workbook()
    hand.active.title = "Wiring"
    hand.active.append(["Fuse", "7.5 A"])
    blob = io.BytesIO()
    hand.save(blob)
    assert keys(data.load(blob.getvalue()).rows()) == ["Fuse"]


# ------------------------------------------------------------------ not trampling a person's file


def test_a_hand_edited_workbook_keeps_its_formatting_and_its_extra_column():
    hand = Workbook()
    sheet = hand.active
    sheet.title = "Torque specs"
    sheet.append(["Key", "Value", "Note", "Checked by"])
    sheet.append(["Caliper bolt torque", "25 Nm", "dry thread", "Marco"])
    sheet["A1"].font = Font(bold=True, italic=True)
    sheet.column_dimensions["A"].width = 61
    first = io.BytesIO()
    hand.save(first)
    blob = first.getvalue()

    changed = data.load(blob)
    changed.put("Torque specs", "Carrier bolt torque", "110 Nm")
    sheet = load_workbook(io.BytesIO(data.dump(changed, blob)))["Torque specs"]

    assert sheet["A1"].font.italic  # their styling, not ours
    assert sheet.column_dimensions["A"].width == 61
    assert sheet["D2"].value == "Marco"  # a column that is none of our business
    assert sheet["A3"].value == "Carrier bolt torque"


def test_an_empty_sheet_someone_made_is_not_deleted_by_the_next_save():
    hand = Workbook()
    hand.active.title = "Paint"
    blank = io.BytesIO()
    hand.save(blank)
    blob = blank.getvalue()

    changed = data.load(blob)
    changed.put("Dimensions", "Jaw width", "125 mm")
    assert "Paint" in load_workbook(io.BytesIO(data.dump(changed, blob))).sheetnames


def test_a_deleted_row_leaves_no_ghost_behind_it():
    made = book(*VICE)
    blob = data.dump(made)
    made.drop(made.find_key("Carrier bolt torque")[0])
    sheet = load_workbook(io.BytesIO(data.dump(made, blob)))["Torque specs"]
    assert [c.value for c in sheet["A"] if c.value] == ["Key", "Caliper bolt torque"]
    assert keys(data.load(data.dump(made, blob)).rows()) == ["Caliper bolt torque", "Jaw width"]


def test_emptying_a_tab_removes_its_sheet():
    made = book(("Paint", "Cover", "RAL 7016"))
    blob = data.dump(made)
    made.drop(made.rows()[0])
    assert "Paint" not in load_workbook(io.BytesIO(data.dump(made, blob))).sheetnames


# ------------------------------------------------------------------ tabs and keys


def test_a_tab_title_excel_would_reject_is_made_acceptable():
    assert data.sheet_title("Torque/Specs: [metric]") == "Torque Specs metric"
    assert len(data.sheet_title("D" * 60)) == data.MAX_TITLE_CHARS
    assert data.sheet_title("///") == data.FALLBACK_TAB


def test_two_spellings_of_one_tab_do_not_become_two_tabs():
    made = book(("Torque specs", "Caliper bolt", "25 Nm"))
    made.put("torque SPECS", "Carrier bolt", "110 Nm")
    made.put("Torque-specs", "Hub nut", "290 Nm")
    assert made.tabs() == {"Torque specs": 3}


def test_saving_a_key_again_replaces_it_in_place_and_says_so():
    made = book(*VICE)
    assert made.put("Torque specs", "caliper BOLT torque", "28 Nm") is True
    assert made.put("Torque specs", "Hub nut", "290 Nm") is False
    row = made.find_key("Caliper bolt torque")[0]
    assert (row.value, row.note) == ("28 Nm", "dry thread")  # a note kept when none was given
    assert keys(data.load(data.dump(made)).rows())[:2] == [
        "Caliper bolt torque",  # still first: a correction is not a new fact
        "Carrier bolt torque",
    ]


# ------------------------------------------------------------------ finding one again


def test_a_near_miss_through_a_microphone_still_finds_the_row():
    found = data.search(book(*VICE), "caliper torq")
    assert keys(found)[0] == "Caliper bolt torque"


def test_the_key_outranks_the_tab_it_is_filed_under():
    found = data.search(book(*VICE), "jaw width")
    assert keys(found)[0] == "Jaw width"


def test_a_query_about_nothing_written_down_comes_back_empty():
    assert data.search(book(*VICE), "sprocket pitch") == []
    assert data.search(book(*VICE), "") == []


def test_search_returns_at_most_the_limit():
    many = book(*[("Bolts", f"Bolt {n} torque", f"{n} Nm") for n in range(12)])
    assert len(data.search(many, "bolt torque")) == data.DEFAULT_LIMIT


def test_an_exact_key_is_never_ambiguous_with_one_that_resembles_it():
    made = book(("Torque specs", "Bolt", "25 Nm"), ("Torque specs", "Bolt torque", "110 Nm"))
    assert keys(made.find_key("bolt")) == ["Bolt"]


def test_an_ambiguous_key_hands_back_every_candidate_so_the_caller_can_ask():
    assert len(book(*VICE).find_key("bolt torque")) == 2


def test_naming_the_tab_narrows_an_ambiguous_key():
    made = book(("Torque specs", "Front bolt", "25 Nm"), ("Dimensions", "Front bolt", "M10"))
    assert [r.value for r in made.find_key("front bolt", "dimensions")] == ["M10"]
    assert made.find_key("front bolt", "Nonexistent tab") == []
